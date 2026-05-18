"""Replace the existing SMARTER table with a new one matching the 7 official objectives."""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"C:\Users\USER\Desktop\PFE\rapport pfe.docx"

OBJECTIVES_FR = [
    ("Centraliser les données SEO dans une base PostgreSQL unifiée", [
        "Base PostgreSQL unique regroupant GSC, SERP, KPIs et enrichissement NLP",
        "Nombre de tables et de lignes ingérées par source de données",
        "PostgreSQL (Neon) déjà éprouvé en environnement de production",
        "Volumes typiques d'une agence comme WAOO Digital",
        "Livré dès la phase d'ingestion en début de projet",
        "Taux de cohérence des jointures entre les tables ingérées",
        "Ajout de nouvelles sources de données sans refonte du modèle",
    ]),
    ("Automatiser la collecte des données SEO", [
        "Workflows n8n déclenchés automatiquement pour GSC et SERP",
        "Nombre de collectes quotidiennes réussies par projet",
        "Intégration n8n + API Search Console + scraping SERP",
        "Réduction de la dépendance aux consultations manuelles",
        "Collecte quotidienne planifiée",
        "Comparaison avec consultation manuelle sur un échantillon",
        "Fréquence et périmètre configurables, déclenchement manuel possible",
    ]),
    ("Qualifier automatiquement les mots-clés (règles + IA)", [
        "Pipeline règles métier locales + enrichissement Z.AI",
        "Taux d'enrichissement de la table nlp_keyword_enrichment",
        "Pipeline règles + Z.AI éprouvé",
        "Couverture progressive selon les besoins",
        "Intégré au pipeline quotidien",
        "Correspondance avec un échantillon vérifié manuellement",
        "Ajout de nouvelles dimensions de classification possible",
    ]),
    ("Calculer des indicateurs décisionnels SEO", [
        "Sept KPIs formalisés (opportunity_score, CTR Gap, Performance Drift…)",
        "Valeur numérique par mot-clé et par date",
        "Formules issues d'un référentiel public",
        "Calculs déterministes à partir des données ingérées",
        "Recalcul à chaque ingestion",
        "Comparaison avec calcul manuel sur un échantillon de contrôle",
        "Seuils et coefficients configurables",
    ]),
    ("Identifier les opportunités SEO à fort potentiel", [
        "Tri par opportunity_score décroissant",
        "Gain de clics potentiel chiffré pour chaque opportunité",
        "Page Opportunités fonctionnelle basée sur données GSC réelles",
        "Volumes et qualité de données maîtrisés",
        "Rafraîchie à chaque collecte",
        "Correspondance avec les recommandations de l'analyste",
        "Filtres par intention, marque, période",
    ]),
    ("Assurer le suivi continu des mots-clés travaillés", [
        "Indicateur is_tracked et baseline d'optimisation par mot-clé",
        "Évolution de la position et du CTR vs date de référence",
        "Table keywords étendue avec date de début d'optimisation",
        "Volumes maîtrisés par projet",
        "Suivi quotidien",
        "Variation positive observée par rapport à la baseline",
        "Date de référence d'optimisation modifiable à tout moment",
    ]),
    ("Visualiser les performances via tableaux de bord interactifs", [
        "Page Tableau de bord avec KPIs, graphiques et alertes",
        "KPIs lisibles en un coup d'œil",
        "Recharts et React",
        "Composants graphiques éprouvés",
        "Mise à jour quasi temps réel après ingestion",
        "Lisibilité validée auprès de l'analyste",
        "Filtres par projet et par période",
    ]),
]

HEADERS = ["Objectif", "Spécifique", "Mesurable", "Atteignable",
           "Réaliste", "Temporel", "Évaluable", "Réajustable"]


def set_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{edge}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:color'), '000000')
        tcBorders.append(b)
    tcPr.append(tcBorders)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


doc = Document(SRC)

# Locate the existing SMARTER table by its first cell content.
target_table = None
for t in doc.tables:
    try:
        if t.rows[0].cells[0].text.strip() == "Objectif" and t.rows[0].cells[1].text.strip() == "Spécifique":
            target_table = t
            break
    except IndexError:
        continue

if target_table is None:
    raise SystemExit("SMARTER table not found.")

old_el = target_table._element

# Build the new table separately (append at end, then move into place)
new_table = doc.add_table(rows=1 + len(OBJECTIVES_FR), cols=len(HEADERS))
new_table.alignment = WD_TABLE_ALIGNMENT.CENTER
new_table.autofit = True

hdr = new_table.rows[0].cells
for i, h in enumerate(HEADERS):
    hdr[i].text = ""
    p = hdr[i].paragraphs[0]
    r = p.add_run(h)
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    shade_cell(hdr[i], "4472C4")
    set_cell_border(hdr[i])

for ri, (title, crits) in enumerate(OBJECTIVES_FR, start=1):
    row = new_table.rows[ri].cells
    row[0].text = ""
    p0 = row[0].paragraphs[0]
    r0 = p0.add_run(title)
    r0.bold = True
    r0.font.size = Pt(9)
    shade_cell(row[0], "D9E1F2")
    set_cell_border(row[0])
    for ci, val in enumerate(crits, start=1):
        row[ci].text = ""
        pp = row[ci].paragraphs[0]
        rr = pp.add_run(val)
        rr.font.size = Pt(9)
        if ri % 2 == 0:
            shade_cell(row[ci], "F2F2F2")
        set_cell_border(row[ci])

# Move new table just before the old one, then remove the old.
new_el = new_table._element
new_el.getparent().remove(new_el)
old_el.addprevious(new_el)
old_el.getparent().remove(old_el)

doc.save(SRC)
print("SMARTER table replaced with 7 official objectives.")
