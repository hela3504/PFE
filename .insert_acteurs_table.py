"""Replace the actors bullet list with a 2-column table (Acteur / Role)."""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"C:\Users\USER\Desktop\PFE\rapport pfe.docx"

ACTEURS = [
    ("SEO Analyst (Utilisateur)",
     "Acteur principal du système : analyste SEO qui exploite l'application pour analyser les performances SEO d'un site web."),
    ("Orchestrateur d'automatisation (n8n)",
     "Outil d'automatisation qui agit comme orchestrateur du système. Son rôle principal est d'automatiser la collecte des données SEO provenant de différentes sources externes."),
    ("Google Search Console API",
     "Service externe fourni par Google qui expose les données de performance SEO d'un site web (impressions, clics, position, CTR, requêtes, pages)."),
    ("SERP API Provider",
     "Service externe permettant de récupérer des informations sur les résultats des moteurs de recherche pour un mot-clé donné (volume, concurrence, PAA, AI Overview, CPC)."),
    ("API Z.AI",
     "Couche d'intelligence artificielle intégrée à l'application. Elle exploite un modèle de langage de type GLM pour enrichir l'analyse des mots-clés : classification de l'intention de recherche, qualification des opportunités SEO et aide à la priorisation."),
]

HEADERS = ["Acteur", "Rôle"]


def set_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{edge}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:color'), '000000')
        tcBorders.append(b)
    tcPr.append(tcBorders)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


doc = Document(SRC)
paragraphs = list(doc.paragraphs)

# Locate the actor bullets by content (robust to index drift).
start_idx = None
end_idx = None
for i, p in enumerate(paragraphs):
    t = p.text.strip()
    if start_idx is None and t.startswith("SEO Analyst"):
        start_idx = i
    if t.startswith("API Z.AI") or t.startswith("API Z.AI :"):
        end_idx = i

if start_idx is None or end_idx is None:
    raise SystemExit(f"Could not locate actors block (start={start_idx}, end={end_idx})")

print(f"Found actors paragraphs {start_idx}..{end_idx}")
target_paragraphs = paragraphs[start_idx:end_idx + 1]
anchor = target_paragraphs[0]

# Build table
table = doc.add_table(rows=1 + len(ACTEURS), cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = True

# Header
hdr = table.rows[0].cells
for i, h in enumerate(HEADERS):
    hdr[i].text = ""
    p = hdr[i].paragraphs[0]
    r = p.add_run(h)
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    shade_cell(hdr[i], "4472C4")
    set_cell_border(hdr[i])

# Rows
for ri, (name, role) in enumerate(ACTEURS, start=1):
    row = table.rows[ri].cells
    row[0].text = ""
    p0 = row[0].paragraphs[0]
    r0 = p0.add_run(name)
    r0.bold = True
    r0.font.size = Pt(10)
    shade_cell(row[0], "D9E1F2")
    set_cell_border(row[0])

    row[1].text = ""
    p1 = row[1].paragraphs[0]
    r1 = p1.add_run(role)
    r1.font.size = Pt(10)
    if ri % 2 == 0:
        shade_cell(row[1], "F2F2F2")
    set_cell_border(row[1])

# Move the table before the first target paragraph
table_el = table._element
table_el.getparent().remove(table_el)
anchor._element.addprevious(table_el)

# Delete the old bullet paragraphs
for p in target_paragraphs:
    p._element.getparent().remove(p._element)

doc.save(SRC)
print("Saved.")
