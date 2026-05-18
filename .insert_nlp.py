import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.oxml import OxmlElement

PATH = "C:/Users/USER/Desktop/PFE/rapport pfe.docx"
doc = Document(PATH)


def insert_paragraph_after(paragraph, text="", style="Normal", bold=False):
    new_p = OxmlElement('w:p')
    paragraph._p.addnext(new_p)
    new_para = paragraph.__class__(new_p, paragraph._parent)
    if text:
        run = new_para.add_run(text)
        if bold:
            run.bold = True
    return new_para


# Anchor: insert AFTER the last paragraph of the ETL section,
# which ends with "...moteur décisionnel et les tableaux de bord de l'application."
# That's the paragraph right before "1.Environnement de développement"
anchor_idx = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith("1.Environnement de développement") or \
       p.text.strip().startswith("1. Environnement de développement"):
        anchor_idx = i
        break

if anchor_idx is None:
    print("ERROR: Environnement de développement anchor not found")
    sys.exit(1)

# Walk back to last non-empty paragraph (the end of ETL section)
prev_idx = anchor_idx - 1
while prev_idx > 0 and not doc.paragraphs[prev_idx].text.strip():
    prev_idx -= 1

current = doc.paragraphs[prev_idx]
print(f"Inserting after para {prev_idx}: '{current.text[:70]}...'")
print(f"  (before para {anchor_idx}: '{doc.paragraphs[anchor_idx].text[:60]}')")

# Heading 3 style by name
h3 = next(s for s in doc.styles if s.name == "Heading 3")


def add(text, style="Normal", bold=False, heading=False):
    global current
    current = insert_paragraph_after(current, text, style, bold)
    if heading:
        current.style = h3


# --- NLP ENRICHMENT SECTION ---
add("")
add("Qualification NLP des mots-clés", heading=True)
add("")
add(
    "Les indicateurs décisionnels issus du calcul des KPIs (opportunity_score, ctr_gap, etc.) "
    "répondent à la question quantitative « combien de clics peut-on espérer gagner ». Ils ne "
    "renseignent toutefois pas sur la nature qualitative du mot-clé : s'agit-il d'une requête "
    "de marque ou générique ? D'une intention d'achat ou d'une simple recherche d'information ? "
    "Faut-il l'écarter du périmètre d'opportunités, et quelle action concrète recommander à "
    "l'analyste ? Ces questions, traditionnellement traitées manuellement par l'expert SEO, "
    "deviennent insolubles dès que le volume de mots-clés dépasse quelques centaines."
)
add("")
add(
    "La table nlp_keyword_enrichment a été conçue pour répondre à ce besoin. Elle conserve, "
    "pour chaque triplet (projet, mot-clé, date), une qualification multidimensionnelle "
    "produite automatiquement à partir du mot-clé brut. Neuf dimensions sont stockées :"
)
add("")
add(
    "branded_status (« branded » ou « non-branded ») isole les recherches contenant le nom de "
    "marque du projet des requêtes génériques, plus stratégiques car non encore acquises face "
    "à la concurrence.",
    style="List Paragraph"
)
add(
    "search_intent (« informationnelle », « transactionnelle » ou « navigationnelle ») oriente "
    "la nature du contenu à produire : article éditorial, page produit ou landing optimisée.",
    style="List Paragraph"
)
add(
    "stability_status caractérise la régularité de la position du mot-clé.",
    style="List Paragraph"
)
add(
    "tail_type distingue courte, moyenne et longue traîne, conditionnant l'arbitrage entre "
    "volume de recherche et concurrence.",
    style="List Paragraph"
)
add(
    "exclude_from_opportunity est un drapeau booléen qui masque les mots-clés non pertinents "
    "(recherches de marque, typos, bruit) du flux d'opportunités présentées à l'analyste.",
    style="List Paragraph"
)
add(
    "qualification_label fournit une étiquette métier résumant la nature du mot-clé.",
    style="List Paragraph"
)
add(
    "priority_level (« high », « medium » ou « low ») apporte un jugement qualitatif "
    "complémentaire au score numérique.",
    style="List Paragraph"
)
add(
    "action_hint propose une recommandation concrète et exploitable.",
    style="List Paragraph"
)
add(
    "reasoning conserve la justification produite par l'IA, restituée à l'utilisateur dans la "
    "modale de qualification pour expliquer le verdict.",
    style="List Paragraph"
)
add("")
add(
    "Le pipeline d'enrichissement est volontairement hybride afin de maîtriser les coûts "
    "d'appel à l'API Z.AI. Une première étape locale, fondée sur un ensemble de règles "
    "déterministes (correspondance avec les mots-clés de marque déclarés sur le projet, "
    "détection d'intention via des motifs lexicaux tels que « acheter », « prix », « comment », "
    "etc.), classe immédiatement la majorité des mots-clés. Seuls ceux dont la classification "
    "présente un faible niveau de confiance sont délégués à l'API Z.AI, qui retourne une "
    "qualification structurée en JSON. Cette stratégie permet en pratique d'enrichir la "
    "majorité des mots-clés sans appel externe, tout en garantissant une couverture exhaustive "
    "du corpus."
)
add("")
add(
    "L'enrichissement nourrit ensuite plusieurs composants de l'application : les colonnes "
    "« Intention » et « Marque » de la page Opportunités, les badges colorés correspondants, "
    "le filtre d'exclusion des mots-clés de marque, la modale de qualification détaillée, la "
    "répartition par marque affichée sur le tableau de bord, la page Suivi, ainsi que "
    "l'assistant SEO conversationnel lorsque l'utilisateur active le contexte projet. Sans "
    "cette couche de qualification, les mots-clés seraient présentés comme une liste "
    "indifférenciée de paires (texte, score) — exactement la limite que reproche le rapport au "
    "processus manuel actuellement employé chez WAOO Digital."
)
add("")

doc.save(PATH)
print("OK NLP section inserted")
