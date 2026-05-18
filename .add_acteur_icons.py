"""Add an 'Icône' column with emoji icons to the actors table."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"C:\Users\USER\Desktop\PFE\rapport pfe.docx"

ICONS = {
    "SEO Analyst": "👤",
    "Orchestrateur": "⚙️",
    "Google Search Console": "🔍",
    "SERP API": "🌐",
    "API Z.AI": "🤖",
}


def set_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{edge}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:color'), '000000')
        tcBorders.append(b)
    tcPr.append(tcBorders)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def add_emoji_font(run, emoji):
    """Set Segoe UI Emoji font for emoji glyphs so Word renders them in color."""
    run.text = emoji
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    for attr in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
        rFonts.set(qn(f'w:{attr}'), 'Segoe UI Emoji')


doc = Document(SRC)

# Find the actors table by header content
target_table = None
for tbl in doc.tables:
    rows = tbl.rows
    if len(rows) < 2 or len(rows[0].cells) < 2:
        continue
    h0 = rows[0].cells[0].text.strip()
    h1 = rows[0].cells[1].text.strip()
    if h0 == "Acteur" and h1.startswith("R"):
        # Also check second row contains "SEO Analyst"
        if "SEO Analyst" in rows[1].cells[0].text:
            target_table = tbl
            break

if target_table is None:
    raise SystemExit("Actors table not found")

print(f"Found actors table with {len(target_table.rows)} rows, {len(target_table.columns)} cols")

# Skip if already has 3 columns (icon already added)
if len(target_table.columns) >= 3:
    raise SystemExit("Table already has 3+ columns — icon column appears to already exist")

# Insert a new first column for each row by manipulating XML
tbl_el = target_table._element

# Update grid: add a new <w:gridCol> at the start
tblGrid = tbl_el.find(qn('w:tblGrid'))
if tblGrid is not None:
    new_grid = OxmlElement('w:gridCol')
    new_grid.set(qn('w:w'), '900')  # narrow column
    tblGrid.insert(0, new_grid)

# For every row, prepend a new cell
for ri, row in enumerate(target_table.rows):
    row_el = row._element
    # Build new tc
    tc = OxmlElement('w:tc')
    tcPr = OxmlElement('w:tcPr')
    tcW = OxmlElement('w:tcW')
    tcW.set(qn('w:w'), '900')
    tcW.set(qn('w:type'), 'dxa')
    tcPr.append(tcW)
    tc.append(tcPr)
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'center')
    pPr.append(jc)
    p.append(pPr)
    tc.append(p)

    # Insert before existing first <w:tc>
    first_tc = row_el.find(qn('w:tc'))
    row_el.insert(list(row_el).index(first_tc), tc)

# Now refresh row references and populate cells
target_table = None
for tbl in doc.tables:
    rows = tbl.rows
    if len(rows) < 2:
        continue
    if rows[0].cells[1].text.strip() == "Acteur":
        target_table = tbl
        break

if target_table is None:
    raise SystemExit("Couldn't re-locate table after column insertion")

# Header row
hdr_cell = target_table.rows[0].cells[0]
hdr_cell.text = ""
p = hdr_cell.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Icône")
r.bold = True
r.font.size = Pt(10)
r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
shade_cell(hdr_cell, "4472C4")
set_cell_border(hdr_cell)

# Data rows
for ri in range(1, len(target_table.rows)):
    icon_cell = target_table.rows[ri].cells[0]
    actor_name = target_table.rows[ri].cells[1].text
    # pick emoji
    emoji = "•"
    for key, val in ICONS.items():
        if key in actor_name:
            emoji = val
            break

    icon_cell.text = ""
    p = icon_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("")
    r.font.size = Pt(18)
    add_emoji_font(r, emoji)

    # Match the existing row shading (alternating)
    if ri % 2 == 0:
        shade_cell(icon_cell, "F2F2F2")
    else:
        shade_cell(icon_cell, "FFFFFF")
    set_cell_border(icon_cell)

doc.save(SRC)
print("Saved.")
